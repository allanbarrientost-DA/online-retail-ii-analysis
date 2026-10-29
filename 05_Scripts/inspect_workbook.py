from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


# Define the project directories
PROJECT_DIR = Path(__file__).resolve().parent.parent
ORIGINAL_DATA_DIR = PROJECT_DIR / "01_Original_Data"


def find_excel_file():
    """Locate the original Excel workbook."""

    excel_files = list(ORIGINAL_DATA_DIR.rglob("online_retail_II.xlsx"))

    if not excel_files:
        raise FileNotFoundError(
            f"Excel file not found inside: {ORIGINAL_DATA_DIR}"
        )

    if len(excel_files) > 1:
        raise RuntimeError(
            "Multiple files named online_retail_II.xlsx were found. "
            "Please keep only one original copy in the data directory."
        )

    return excel_files[0]


def inspect_workbook(file_path):
    """Display workbook information without modifying the source file."""

    print("=" * 60)
    print("ONLINE RETAIL II - WORKBOOK INSPECTION")
    print("=" * 60)

    print(f"\nFile: {file_path.name}")
    print(f"Location: {file_path}")

    # Open the workbook in read-only mode
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True
    )

    try:
        print("\nWORKSHEET SUMMARY")

        for worksheet in workbook.worksheets:
            print(f"\nWorksheet: {worksheet.title}")
            print(f"Rows including header: {worksheet.max_row}")
            print(f"Columns: {worksheet.max_column}")
            print(f"Estimated data records: {worksheet.max_row - 1}")

        # Display the first five rows of each worksheet
        print("\nFIRST FIVE ROWS")

        for worksheet_name in workbook.sheetnames:
            print(f"\n{'-' * 60}")
            print(f"Worksheet: {worksheet_name}")
            print("-" * 60)

            preview = pd.read_excel(
                file_path,
                sheet_name=worksheet_name,
                nrows=5
            )

            print("\nColumn names:")
            print(preview.columns.tolist())

            print("\nPreview:")
            print(preview.to_string(index=False))

    finally:
        workbook.close()

    print("\nInspection completed successfully.")


if __name__ == "__main__":
    excel_file = find_excel_file()
    inspect_workbook(excel_file)